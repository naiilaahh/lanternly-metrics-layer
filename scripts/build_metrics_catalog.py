"""
Generate the published metrics catalog from the governed YAML.

    models/semantic/_metrics.yml          (definitions + governance metadata)  ─┐
    models/semantic/_semantic_models.yml  (measures -> dbt models, lineage)     ├─> metrics_catalog/metrics_catalog.csv
    docs/metric_logic_rules.md            (edge-case rules)                    ─┘   metrics_catalog/metrics_catalog.xlsx
                                                                                    docs/metrics_catalog.md

The catalog is never edited by hand. CI regenerates it and fails if the committed copy
is stale (`git diff --exit-code metrics_catalog/metrics_catalog.csv docs/metrics_catalog.md`;
the .xlsx is rebuilt on release because Excel files are not byte-reproducible).

Usage: python scripts/build_metrics_catalog.py
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "metrics_catalog"

sem = yaml.safe_load(open(ROOT / "models/semantic/_semantic_models.yml"))["semantic_models"]
metrics = yaml.safe_load(open(ROOT / "models/semantic/_metrics.yml"))["metrics"]
by_name = {m["name"]: m for m in metrics}
measure_home = {}
for sm in sem:
    model = re.search(r"ref\('([^']+)'\)", sm["model"]).group(1)
    for ms in sm.get("measures", []):
        measure_home[ms["name"]] = (sm["name"], model)

rules_md = open(ROOT / "docs/metric_logic_rules.md").read()
rule_rows = re.findall(r"^\| (MLR-\d{3}) \| (.+?) \| (.+?) \| (.+?) \|$", rules_md, re.M)


def gov(m):
    return m["config"]["meta"]["governance"]


def leaf_measures(name: str, seen=None) -> list[str]:
    seen = seen or set()
    m = by_name[name]
    tp = m["type_params"]
    if m["type"] in ("simple", "cumulative"):
        return [tp["measure"]]
    deps = [tp["numerator"], tp["denominator"]] if m["type"] == "ratio" else [x["name"] for x in tp["metrics"]]
    out = []
    for x in deps:
        if x not in seen:
            seen.add(x)
            out += leaf_measures(x, seen)
    return list(dict.fromkeys(out))


def lineage(name: str) -> tuple[str, str]:
    ms = leaf_measures(name)
    sms = list(dict.fromkeys(measure_home[x][0] for x in ms))
    models = list(dict.fromkeys(measure_home[x][1] for x in ms))
    return ", ".join(sms), ", ".join(models)


TIER_LABEL = {1: "Tier 1 · Board", 2: "Tier 2 · Operating", 3: "Tier 3 · Component"}
DOMAIN_ORDER = ["Recurring revenue", "Customers & retention", "Revenue & margin", "Unit economics", "Product engagement"]
rows = []
for m in sorted(metrics, key=lambda m: (DOMAIN_ORDER.index(gov(m)["domain"]), gov(m)["tier"], m["name"])):
    g = gov(m)
    sms, models = lineage(m["name"])
    rows.append({
        "metric": m["name"],
        "label": m["label"],
        "domain": g["domain"],
        "tier": TIER_LABEL[g["tier"]],
        "status": g["status"],
        "version": g["version"],
        "definition": " ".join(m["description"].split()),
        "formula": g["formula"],
        "metric_type": m["type"],
        "business_owner": g["business_owner"],
        "data_steward": g["data_steward"],
        "refresh": g["refresh"],
        "logic_rules": ", ".join(g.get("logic_rules") or []),
        "semantic_models": sms,
        "dbt_models": models,
        "consumers": "; ".join(g["consumers"]),
        "certified_on": g.get("certified_on") or "",
        "next_review": g.get("next_review") or "",
        "deprecated_on": g.get("deprecated_on") or "",
        "sunset_on": g.get("sunset_on") or "",
        "replaced_by": g.get("replaced_by") or "",
        "known_limitations": g.get("known_limitations") or "",
    })

# ---------------------------------------------------------------- CSV
OUT.mkdir(exist_ok=True)
with open(OUT / "metrics_catalog.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
    w.writeheader()
    w.writerows(rows)

# ---------------------------------------------------------------- XLSX
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
THIN = Side(style="thin", color="D0D7DE")
STATUS_FILL = {"certified": "DFF3E4", "beta": "FFF4CC", "deprecated": "FBE3E3", "proposed": "E8EEF7"}


def sheet(ws, header, data, widths, table_name):
    ws.append(header)
    for r in data:
        ws.append(r)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=THIN)
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "C2" if len(header) > 6 else "A2"
    ws.row_dimensions[1].height = 30
    ref = f"A1:{get_column_letter(len(header))}{len(data) + 1}"
    t = Table(displayName=table_name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)


wb = Workbook()

# Summary (formulas, so it stays live if someone filters/edits the catalog tab)
ws = wb.active
ws.title = "Summary"
ws["A1"] = "Lanternly Metrics Catalog"
ws["A1"].font = Font(name=FONT, bold=True, size=16, color="1F3A5F")
ws["A2"] = ("Generated from models/semantic/_metrics.yml by scripts/build_metrics_catalog.py. "
            "Do not edit by hand: change the YAML through a metric change request (RFC-0001).")
ws["A2"].font = Font(name=FONT, italic=True, size=9, color="555555")
ws["A4"], ws["B4"] = "Status", "Metrics"
for c in (ws["A4"], ws["B4"]):
    c.fill, c.font = HDR_FILL, HDR_FONT
statuses = ["certified", "beta", "deprecated"]
for i, s in enumerate(statuses, start=5):
    ws.cell(i, 1, s.capitalize()).font = BODY
    ws.cell(i, 2, f'=COUNTIF(Catalog!E:E,"{s}")').font = BODY
tot = 5 + len(statuses)
ws.cell(tot, 1, "Total").font = Font(name=FONT, bold=True, size=10)
ws.cell(tot, 2, f"=SUM(B5:B{tot - 1})").font = Font(name=FONT, bold=True, size=10)

ws.cell(tot + 2, 1, "Tier").fill = HDR_FILL
ws.cell(tot + 2, 2, "Metrics").fill = HDR_FILL
ws.cell(tot + 2, 3, "Of which certified").fill = HDR_FILL
for c in range(1, 4):
    ws.cell(tot + 2, c).font = HDR_FONT
for i, t in enumerate(TIER_LABEL.values(), start=tot + 3):
    ws.cell(i, 1, t).font = BODY
    ws.cell(i, 2, f'=COUNTIF(Catalog!D:D,"{t}")').font = BODY
    ws.cell(i, 3, f'=COUNTIFS(Catalog!D:D,"{t}",Catalog!E:E,"certified")').font = BODY

r0 = tot + 3 + len(TIER_LABEL) + 1
ws.cell(r0, 1, "Domain").fill = HDR_FILL
ws.cell(r0, 2, "Metrics").fill = HDR_FILL
ws.cell(r0, 1).font = ws.cell(r0, 2).font = HDR_FONT
for i, d in enumerate(DOMAIN_ORDER, start=r0 + 1):
    ws.cell(i, 1, d).font = BODY
    ws.cell(i, 2, f'=COUNTIF(Catalog!C:C,"{d}")').font = BODY
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 18

# Catalog
ws = wb.create_sheet("Catalog")
hdr = ["Metric", "Label", "Domain", "Tier", "Status", "Version", "Business definition", "Formula", "Type",
       "Business owner", "Data steward", "Refresh / SLA", "Logic rules", "Semantic models", "dbt models",
       "Consumers", "Certified on", "Next review", "Deprecated on", "Sunset on", "Replaced by", "Known limitations"]
sheet(ws, hdr, [list(r.values()) for r in rows],
      [22, 28, 18, 18, 11, 9, 60, 44, 10, 24, 26, 30, 22, 28, 30, 36, 12, 12, 13, 12, 14, 40], "Catalog")
for row in ws.iter_rows(min_row=2):
    st = row[4].value
    row[4].fill = PatternFill("solid", fgColor=STATUS_FILL.get(st, "FFFFFF"))

# Governance matrix (RACI per domain)
ws = wb.create_sheet("Governance Matrix")
raci = []
for d in DOMAIN_ORDER:
    owners = sorted({r["business_owner"] for r in rows if r["domain"] == d})
    stewards = sorted({r["data_steward"] for r in rows if r["domain"] == d})
    t1 = sum(1 for r in rows if r["domain"] == d and r["tier"].startswith("Tier 1"))
    raci.append([d, "; ".join(owners), "; ".join(stewards), "Metrics Council (tier-1 changes)",
                 "BI developers, analysts, embedded-app team", t1,
                 "Council + business owner + steward" if t1 else "Business owner + steward"])
sheet(ws, ["Domain", "Accountable (business owner)", "Responsible (data steward)", "Consulted",
           "Informed", "Tier-1 metrics", "Approval needed for logic change"], raci,
      [22, 40, 34, 30, 36, 12, 34], "GovernanceMatrix")

# Change-approval matrix
ws = wb.create_sheet("Change Approvals")
sheet(ws, ["Change type", "Example", "Version bump", "Tier 1 approvers", "Tier 2 approvers", "Tier 3 approvers",
           "Notice to consumers"], [
    ["Logic change that alters history", "Exclude Customer Success from CAC", "MAJOR",
     "Metrics Council + owner + steward", "Owner + steward", "Steward", "30 days, dual-run 1 close"],
    ["New dimension or filter", "Slice NRR by region", "MINOR", "Owner + steward", "Steward", "Steward", "Release notes"],
    ["Docs / description only", "Clarify ARR excludes usage", "PATCH", "Steward", "Steward", "Steward", "None"],
    ["New metric", "Product-qualified accounts", "0.1.0 (beta)", "Council review to certify", "Owner + steward",
     "Steward", "Catalog entry"],
    ["Deprecation", "Retire total_revenue_legacy", "-", "Metrics Council", "Owner", "Steward",
     "90 days minimum before sunset"],
], [30, 34, 13, 32, 22, 16, 28], "ChangeApprovals")

# Logic rules
ws = wb.create_sheet("Logic Rules")
sheet(ws, ["Rule ID", "Rule", "Implemented in", "Tested by"],
      [[a, b.replace("`", ""), c.replace("`", ""), d.replace("`", "")] for a, b, c, d in rule_rows],
      [10, 70, 50, 60], "LogicRules")

# Lineage
ws = wb.create_sheet("Lineage")
lin = []
for r in rows:
    for ms in leaf_measures(r["metric"]):
        lin.append([r["metric"], ms, measure_home[ms][0], measure_home[ms][1]])
sheet(ws, ["Metric", "Measure", "Semantic model", "dbt model"], lin, [24, 34, 22, 30], "Lineage")

wb.save(OUT / "metrics_catalog.xlsx")

# ---------------------------------------------------------------- Markdown
md = ["# Metrics Catalog", "",
      "> Generated by `scripts/build_metrics_catalog.py` from `models/semantic/_metrics.yml`. "
      "Do not edit by hand. Change the YAML through a metric change request ([RFC-0001](rfc/RFC-0001-metrics-governance.md)).",
      "",
      f"**{len(rows)} metrics**: "
      + ", ".join(f"{sum(r['status'] == s for r in rows)} {s}" for s in statuses)
      + ". Edge-case rules referenced below are defined in [metric_logic_rules.md](metric_logic_rules.md).", ""]
BADGE = {"certified": "✅ certified", "beta": "🧪 beta", "deprecated": "⛔ deprecated"}
for d in DOMAIN_ORDER:
    md += [f"## {d}", "", "| Metric | Tier | Status | Definition | Owner | Rules |", "|---|---|---|---|---|---|"]
    for r in [r for r in rows if r["domain"] == d]:
        extra = f" **Replaced by `{r['replaced_by']}` on {r['sunset_on']}.**" if r["replaced_by"] else ""
        md.append(f"| `{r['metric']}`<br>{r['label']} | {r['tier'].split(' · ')[0]} | {BADGE[r['status']]} v{r['version']} "
                  f"| {r['definition']}{extra}<br><sub>Formula: {r['formula']}</sub> | {r['business_owner']} | {r['logic_rules']} |")
    md.append("")
md += ["## Lineage", "", "| Metric | Semantic models | dbt models |", "|---|---|---|"]
md += [f"| `{r['metric']}` | {r['semantic_models']} | {r['dbt_models']} |" for r in rows]
(ROOT / "docs/metrics_catalog.md").write_text("\n".join(md) + "\n")

print(f"Catalog built: {len(rows)} metrics -> metrics_catalog/metrics_catalog.csv, .xlsx, docs/metrics_catalog.md")
